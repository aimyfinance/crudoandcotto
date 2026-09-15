"""Експорт: Excel у структурі звіту менеджера (аркуші «Товар», «Зведення», «Оплати»,
«Операції», «Залишки») та CSV операцій/залишків."""
from __future__ import annotations

import csv
import io
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import services as S
from .db import Database, local_dt_str, local_date_of
from .money import d, round_cents

HDR_FILL = PatternFill("solid", fgColor="DDEBF7")
BOLD = Font(bold=True)


def _f(v):
    """Decimal -> float для Excel (значення вже округлені до центів)."""
    if v is None:
        return None
    return float(v)


def _kg(g):
    return round(g / 1000, 3) if g is not None else None


def _autosize(ws):
    for col in ws.columns:
        width = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(10, width + 2), 45)


def _header(ws, row, values, bold=True):
    for i, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=i, value=v)
        if bold:
            cell.font = BOLD
            cell.fill = HDR_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_excel(db: Database, date_from: str, date_to: str, out_path: Path) -> Path:
    wb = Workbook()

    # ---------- Аркуш «Товар» — структура як у файлі менеджера ----------
    ws = wb.active
    ws.title = "Товар"
    ws.cell(row=1, column=1, value=f"Товари {date_from} — {date_to}").font = BOLD
    _header(ws, 2, ["Категорія", "Найменування", "Партія", "Дата", "Придатний до",
                    "Залишок на початок, кг", "Ціна, €/кг", "Сума, €",
                    "Купівля, кг", "Ціна постач., €/кг", "Собівартість, €/кг", "Сума, €",
                    "Продаж, кг", "Сер. ціна, €/кг", "Сума, €", "Прибуток, €",
                    "Втрати, кг", "Сума, €",
                    "Залишок на кінець, кг", "Ціна, €/кг", "Сума, €"])
    r = 3
    rows = S.product_ledger(db, date_from, date_to)
    current = None
    tot = {"opening_c": Decimal(0), "buy_c": Decimal(0), "sold_rev": Decimal(0), "profit": Decimal(0),
           "loss_c": Decimal(0), "closing_c": Decimal(0)}
    for row in rows:
        if current and row["product"] != current:
            r += 1
        current = row["product"]
        vals = [row["category"], row["product"], row["batch"], row["date"], row["expiry"],
                _kg(row["opening_g"]), _f(row["landed_price"]), _f(row["opening_c"]),
                _kg(row["buy_g"]), _f(row["buy_price"]), _f(row["landed_price"]), _f(row["buy_c"]),
                _kg(row["sold_g"]), _f(round_cents(row["sold_price"])) if row["sold_price"] else None, _f(row["sold_rev"]),
                _f(row["profit"]),
                _kg(row["loss_g"]), _f(row["loss_c"]),
                _kg(row["closing_g"]), _f(row["landed_price"]), _f(row["closing_c"])]
        for i, v in enumerate(vals, 1):
            ws.cell(row=r, column=i, value=v)
        for k in tot:
            tot[k] += row[k]
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Всього").font = BOLD
    for col, key in ((8, "opening_c"), (12, "buy_c"), (15, "sold_rev"), (16, "profit"), (18, "loss_c"), (21, "closing_c")):
        ws.cell(row=r, column=col, value=_f(tot[key])).font = BOLD
    ws.freeze_panes = "C3"
    _autosize(ws)

    # ---------- Аркуш «Зведення» — як «Лист3» ----------
    ws2 = wb.create_sheet("Зведення")
    ws2.cell(row=1, column=1, value=f"{date_from} — {date_to}").font = BOLD
    _header(ws2, 2, ["Товари", "Залишок на початок, кг", "євро", "Надходження, кг", "євро",
                     "Витрати (продаж+списання), кг", "євро (собівартість)", "Виручка, €", "Прибуток, €",
                     "Залишок на кінець, кг", "євро"])
    agg: dict[str, dict] = {}
    for row in rows:
        a = agg.setdefault(row["product"], {k: Decimal(0) for k in
                                            ("og", "oc", "bg", "bc", "ug", "uc", "rev", "pf", "cg", "cc")})
        a["og"] += row["opening_g"]; a["oc"] += row["opening_c"]
        a["bg"] += row["buy_g"]; a["bc"] += row["buy_c"]
        a["ug"] += row["sold_g"] + row["loss_g"]; a["uc"] += row["sold_cost"] + row["loss_c"]
        a["rev"] += row["sold_rev"]; a["pf"] += row["profit"]
        a["cg"] += row["closing_g"]; a["cc"] += row["closing_c"]
    r = 3
    for name, a in agg.items():
        vals = [name, _kg(int(a["og"])), _f(a["oc"]), _kg(int(a["bg"])), _f(a["bc"]), _kg(int(a["ug"])), _f(a["uc"]),
                _f(a["rev"]), _f(a["pf"]), _kg(int(a["cg"])), _f(a["cc"])]
        for i, v in enumerate(vals, 1):
            ws2.cell(row=r, column=i, value=v)
        r += 1
    r += 1
    ws2.cell(row=r, column=1, value="Всього").font = BOLD
    for col, key in ((3, "oc"), (5, "bc"), (7, "uc"), (8, "rev"), (9, "pf"), (11, "cc")):
        ws2.cell(row=r, column=col, value=_f(sum((a[key] for a in agg.values()), Decimal(0)))).font = BOLD
    rep = S.report_period(db, date_from, date_to)
    r += 1
    ws2.cell(row=r, column=1, value="в т.ч. продаж (виручка)"); ws2.cell(row=r, column=8, value=_f(rep["revenue"]))
    r += 1
    ws2.cell(row=r, column=1, value="в т.ч. списання (собівартість)"); ws2.cell(row=r, column=7, value=_f(rep["writeoff_cost"]))
    r += 1
    ws2.cell(row=r, column=1, value="Валовий прибуток (виручка − собівартість проданого)"); ws2.cell(row=r, column=9, value=_f(rep["gross_profit"]))
    r += 1
    ws2.cell(row=r, column=1, value="Додаткові закупівельні витрати (вже включені в собівартість)"); ws2.cell(row=r, column=5, value=_f(rep["purchase_extra_costs"]))
    _autosize(ws2)

    # ---------- Аркуш «Оплати» — для аркуша «Витрати»: готівка / картка по днях ----------
    ws3 = wb.create_sheet("Оплати")
    _header(ws3, 1, ["Дата", "Готівка (бот), €", "Картка (бот), €", "Разом (бот), €",
                     "Готівка (каса), €", "Картка (каса), €", "Разом (каса), €", "Розбіжність, €"])
    r = 2
    for e in S.reconcile(db, date_from, date_to):
        vals = [e["day"], _f(e["bot_cash"]), _f(e["bot_card"]), _f(e["bot_total"]),
                _f(e["reg_cash"]), _f(e["reg_card"]), _f(e["reg_total"]), _f(e["diff"])]
        for i, v in enumerate(vals, 1):
            ws3.cell(row=r, column=i, value=v)
        r += 1
    _autosize(ws3)

    # ---------- Аркуш «Продажі за товарами» ----------
    ws4 = wb.create_sheet("Продажі за товарами")
    _header(ws4, 1, ["Товар", "Категорія", "Продано, кг", "Шт", "Виручка, €", "Собівартість, €", "Валовий прибуток, €", "Маржа, %"])
    for i, e in enumerate(rep["by_product"], 2):
        for j, v in enumerate([e["name"], S.CATEGORIES[e["category"]], _kg(e["grams"]), e["pieces"] or None,
                               _f(e["amount"]), _f(e["cost"]), _f(e["gross_profit"]), round(float(e["margin_pct"]), 1)], 1):
            ws4.cell(row=i, column=j, value=v)
    _autosize(ws4)

    # ---------- Аркуш «Залишки» ----------
    ws5 = wb.create_sheet("Залишки")
    _header(ws5, 1, ["Товар", "Категорія", "Партія", "Дата надходження", "Придатний до", "Залишок, кг",
                     "Собівартість, €/кг", "Вартість залишку, €"])
    r = 2
    for st in S.stock_summary(db):
        for b in S.batches_of_product(db, st["product"]["id"]):
            vals = [st["product"]["name"], S.CATEGORIES[st["product"]["category"]], b["batch_code"] or f"#{b['id']}",
                    b["received_at"][:10], b["expiry_date"], _kg(b["grams_left"]), _f(d(b["landed_price_per_kg"])),
                    _f(round_cents(Decimal(b["grams_left"]) * d(b["landed_price_per_kg"]) / 1000))]
            for i, v in enumerate(vals, 1):
                ws5.cell(row=r, column=i, value=v)
            r += 1
    _autosize(ws5)

    # ---------- Аркуш «Операції» — журнал усіх рухів ----------
    ws6 = wb.create_sheet("Операції")
    _header(ws6, 1, ["Дата/час (Відень)", "Тип", "Документ", "Товар", "Партія", "Δ вага, кг", "Δ собівартість, €", "Користувач"])
    r = 2
    for m in _movements(db, date_from, date_to):
        for i, v in enumerate([local_dt_str(m["ts"]), KIND_UA.get(m["kind"], m["kind"]), f"{m['ref_type']} #{m['ref_id']}",
                               m["product_name"], m["batch_code"] or f"#{m['batch_id']}", _kg(m["grams_delta"]),
                               _f(d(m["cost_delta"])), m["user_name"] or m["user_id"]], 1):
            ws6.cell(row=r, column=i, value=v)
        r += 1
    _autosize(ws6)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


KIND_UA = {
    "purchase": "Закупівля", "opening": "Початковий залишок", "sale": "Продаж", "sale_cancel": "Скасування продажу",
    "writeoff": "Списання", "writeoff_cancel": "Скасування списання", "adjustment": "Інвентаризація",
    "adjustment_cancel": "Скасування інвентаризації", "purchase_cancel": "Скасування закупівлі",
}


def _movements(db: Database, date_from: str, date_to: str):
    # беремо із запасом і фільтруємо за датою Відня
    rows = db.q(
        "SELECT m.*, p.name AS product_name, b.batch_code, u.name AS user_name "
        "FROM stock_movements m JOIN products p ON p.id=m.product_id JOIN batches b ON b.id=m.batch_id "
        "LEFT JOIN users u ON u.telegram_id=m.user_id WHERE m.ts >= ? AND m.ts <= ? ORDER BY m.ts, m.id",
        (f"{date_from}T00:00:00", f"{date_to}T23:59:59+99"),
    )
    # ts у форматі ISO зі зміщенням; '+99' у верхній межі = «все за цю дату»
    return [m for m in rows if date_from <= local_date_of(m["ts"]) <= date_to]


def build_csv_movements(db: Database, date_from: str, date_to: str) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(["datetime_vienna", "kind", "ref", "product", "batch", "grams_delta", "cost_delta_eur", "user"])
    for m in _movements(db, date_from, date_to):
        w.writerow([local_dt_str(m["ts"]), m["kind"], f"{m['ref_type']}#{m['ref_id']}", m["product_name"],
                    m["batch_code"] or m["batch_id"], m["grams_delta"], str(d(m["cost_delta"])).replace(".", ","),
                    m["user_name"] or m["user_id"]])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def build_csv_stock(db: Database) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(["product", "category", "batch", "received", "expiry", "grams_left", "landed_price_per_kg", "value_eur"])
    for st in S.stock_summary(db):
        for b in S.batches_of_product(db, st["product"]["id"]):
            w.writerow([st["product"]["name"], st["product"]["category"], b["batch_code"] or b["id"], b["received_at"][:10],
                        b["expiry_date"] or "", b["grams_left"], str(d(b["landed_price_per_kg"])).replace(".", ","),
                        str(round_cents(Decimal(b["grams_left"]) * d(b["landed_price_per_kg"]) / 1000)).replace(".", ",")])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")
