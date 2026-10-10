from decimal import Decimal
from openpyxl import Workbook
from bot import services as S
from bot.db import Database, set_db
from bot.inventory_xlsx import import_inventory_xlsx


def test_inventory_import(tmp_path):
    db = Database(tmp_path / "t.db"); set_db(db); S.ensure_admins(db, [1])
    tal = S.create_product(db, "Taleggio DOP VERO ROGNONI", "cheese", "weight", Decimal("29"))
    S.add_opening_stock(db, 1, tal, 6692, Decimal("8.25"), None, "2026-10-09", "поставка")
    spi = S.create_product(db, "Spianata Calabrese piccante", "meat", "weight", Decimal("39"))
    S.add_opening_stock(db, 1, spi, 620, Decimal("0"), None, "2026-10-09", "продано за касою, залишку в боті не було")
    S.inventory_adjust(db, 1, S.batches_of_product(db, spi)[0]["id"], 0, "продано")  # ніби продали
    roc = S.create_product(db, "Roccetta", "cheese", "weight", Decimal("39"))
    S.add_opening_stock(db, 1, roc, 336, Decimal("0"), None, "2026-10-10", "початковий")
    fav = S.create_product(db, "Mortadella Favola Gran Reserva", "meat", "weight", Decimal("38"))
    bol = S.create_product(db, "Mortadella Bologna IGP con Pistacchio", "meat", "weight", Decimal("38"))
    S.set_alias(db, "Mortadella Gran Reserva", fav, 1)

    wb = Workbook(); ws = wb.active; ws.title = "Інвентаризація"
    ws.append(["Товар", "У боті", "Очікувано", "Факт, кг", "Ціна, €/кг", "Коментар"])
    ws.append(["Taleggio DOP VERO ROGNONI", 6.69, 4.77, 4.75, None, "зважено"])
    ws.append(["Roccetta", 0.336, 0.76, 0.8, 14.6, ""])
    ws.append(["Chevre Buche", 1, 1, 1, None, ""])          # нема такого товару → помилка
    ws2 = wb.create_sheet("Закупівля"); ws2.append(["Дата", "Постачальник", "Товар", "кг", "€/кг", "Лот", "Термін"])
    ws2.append(["17.09.2026", "Rognoni", "Spianata Calabrese piccante", 5.672, 9.6, "309/26", "17.02.2027"])
    ws3 = wb.create_sheet("Прив'язки каси"); ws3.append(["Назва в касі", "Товар бота"])
    ws3.append(["Mortadella Gran Reserva", "Mortadella Bologna IGP con Pistacchio"])
    p = tmp_path / "inv.xlsx"; wb.save(p)

    res = import_inventory_xlsx(1, p)
    assert res["purchases"] == 1 and res["aliases"] == 1
    assert any("Chevre" in e for e in res["errors"]) and len(res["errors"]) == 1
    assert S.stock_of_product(db, tal) == 4750
    assert S.stock_of_product(db, spi) == 5672 - 620           # закупівля поглинула продане наперед
    assert S.stock_of_product(db, roc) == 800
    assert all(Decimal(b["landed_price_per_kg"]) == Decimal("14.6") for b in S.batches_of_product(db, roc))
    assert S.match_product(db, "Mortadella Gran Reserva") == (bol, "alias")
