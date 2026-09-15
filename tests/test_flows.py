"""Наскрізні тести сценаріїв через Dispatcher з підробленою сесією Telegram."""
import datetime as dt
from decimal import Decimal

import pytest
from aiogram import Bot
from aiogram.client.default import DefaultBotProperties
from aiogram.client.session.base import BaseSession
from aiogram.fsm.storage.base import StorageKey
from aiogram.methods import TelegramMethod
from aiogram.types import CallbackQuery, Chat, Message, Update, User

from bot import services as S
from bot.db import Database, set_db
from bot.keyboards import M_PURCHASE, M_SALE, TODAY, SKIP
from bot.main import build_dispatcher

ADMIN = 1001
SELLER = 1002
STRANGER = 1003


class FakeSession(BaseSession):
    def __init__(self):
        super().__init__()
        self.calls: list[TelegramMethod] = []
        self._mid = 100

    async def close(self):
        pass

    async def stream_content(self, *a, **k):  # pragma: no cover
        yield b""

    async def make_request(self, bot, method, timeout=None):
        self.calls.append(method)
        name = type(method).__name__
        if name in ("SendMessage", "SendDocument"):
            self._mid += 1
            rm = getattr(method, "reply_markup", None)
            return Message(message_id=self._mid, date=dt.datetime.now(), chat=Chat(id=method.chat_id, type="private"),
                           text=getattr(method, "text", None),
                           reply_markup=rm if rm is not None and hasattr(rm, "inline_keyboard") else None)
        return True

    def texts(self):
        return [c.text for c in self.calls if type(c).__name__ == "SendMessage"]

    def last_inline(self):
        for c in reversed(self.calls):
            rm = getattr(c, "reply_markup", None)
            if rm is not None and hasattr(rm, "inline_keyboard"):
                return rm
        return None


@pytest.fixture(scope="module")
def dispatcher():
    return build_dispatcher()   # роутери — модульні синглтони, будуємо один раз


@pytest.fixture
def env(tmp_path, dispatcher):
    db = Database(tmp_path / "flow.db")
    set_db(db)
    S.ensure_admins(db, [ADMIN])
    S.upsert_user(db, SELLER, "seller", "Марія")
    session = FakeSession()
    bot = Bot("123:TEST", session=session, default=DefaultBotProperties(parse_mode="HTML"))
    dispatcher.storage.storage.clear()
    return db, bot, dispatcher, session


_uid = [0]
_mid = [0]


def _user(tg_id):
    return User(id=tg_id, is_bot=False, first_name="U", last_name=str(tg_id))


async def send(dp, bot, tg_id, text):
    _uid[0] += 1
    _mid[0] += 1
    m = Message(message_id=_mid[0], date=dt.datetime.now(), chat=Chat(id=tg_id, type="private"), from_user=_user(tg_id), text=text)
    await dp.feed_update(bot, Update(update_id=_uid[0], message=m))


async def click(dp, bot, tg_id, data, update_id=None):
    if update_id is None:
        _uid[0] += 1
        update_id = _uid[0]
    _mid[0] += 1
    m = Message(message_id=_mid[0], date=dt.datetime.now(), chat=Chat(id=tg_id, type="private"), from_user=_user(tg_id), text="x")
    cb = CallbackQuery(id=f"cb{update_id}", from_user=_user(tg_id), chat_instance="ci", message=m, data=data)
    await dp.feed_update(bot, Update(update_id=update_id, callback_query=cb))


async def state_data(dp, bot, tg_id):
    ctx = dp.fsm.get_context(bot, chat_id=tg_id, user_id=tg_id)
    return await ctx.get_data()


@pytest.mark.asyncio
async def test_unauthorized_user_blocked(env):
    db, bot, dp, s = env
    await send(dp, bot, STRANGER, "/start")
    assert any("Доступ заборонено" in t for t in s.texts())
    assert str(STRANGER) in s.texts()[-1]


@pytest.mark.asyncio
async def test_purchase_then_sale_flow_and_double_confirm(env):
    db, bot, dp, s = env
    pid = S.create_product(db, "Mortadella", "meat", "weight", Decimal("36.30"))

    # --- закупівля менеджером/адміном ---
    await send(dp, bot, ADMIN, M_PURCHASE)
    await send(dp, bot, ADMIN, TODAY)
    await send(dp, bot, ADMIN, "Salumificio Rossi")
    await click(dp, bot, ADMIN, f"pp:pur:id:{pid}")
    await send(dp, bot, ADMIN, "15 кг")
    await send(dp, bot, ADMIN, "23,55")
    await send(dp, bot, ADMIN, "30.11.2026")
    await send(dp, bot, ADMIN, "L-77")
    await click(dp, bot, ADMIN, "pur:extra")
    await send(dp, bot, ADMIN, "15")           # 15 € транспорту на 15 кг = +1 €/кг
    await click(dp, bot, ADMIN, "pur:confirm")
    assert S.stock_of_product(db, pid) == 15000
    b = S.batches_of_product(db, pid)[0]
    assert b["landed_price_per_kg"] == "24.5500" and b["batch_code"] == "L-77"
    assert any("Закупівлю №1 проведено" in t for t in s.texts())

    # --- продаж продавцем ---
    await send(dp, bot, SELLER, M_SALE)
    await click(dp, bot, SELLER, f"pp:sale:id:{pid}")
    await send(dp, bot, SELLER, "250")
    assert any("= <b>9,08 €</b>" in t for t in s.texts())   # 0.25 × 36.30 = 9.075 -> 9.08
    await click(dp, bot, SELLER, "sale:add")
    await click(dp, bot, SELLER, "sale:more")
    await click(dp, bot, SELLER, f"pp:sale:id:{pid}")
    await send(dp, bot, SELLER, "0,5")
    await click(dp, bot, SELLER, "sale:chprice")
    await send(dp, bot, SELLER, "30")            # знижена ціна для цієї позиції
    await click(dp, bot, SELLER, "sale:add")
    data = await state_data(dp, bot, SELLER)
    assert len(data["cart"]) == 2
    key = data["key"]
    await click(dp, bot, SELLER, f"sale:pay:card:{key}")
    sale = db.one("SELECT * FROM sales")
    assert sale["total"] == "24.08" and sale["payment_method"] == "card"
    assert S.stock_of_product(db, pid) == 14250

    # повторне натискання тієї ж кнопки — не дублює продаж
    await click(dp, bot, SELLER, f"sale:pay:card:{key}")
    assert db.one("SELECT COUNT(*) FROM sales")[0] == 1
    assert S.stock_of_product(db, pid) == 14250

    # повторна доставка того самого update_id — ігнорується
    used = _uid[0]
    await click(dp, bot, SELLER, f"pp:sale:id:{pid}", update_id=used)
    assert db.one("SELECT COUNT(*) FROM sales")[0] == 1


@pytest.mark.asyncio
async def test_sale_over_stock_rejected_and_back_button(env):
    db, bot, dp, s = env
    pid = S.create_product(db, "Taleggio", "cheese", "weight", Decimal("28"))
    S.create_purchase(db, ADMIN, "2026-09-01", "A", [S.PurchaseLine(pid, 300, Decimal("19.70"))])
    await send(dp, bot, SELLER, M_SALE)
    await click(dp, bot, SELLER, f"pp:sale:id:{pid}")
    await send(dp, bot, SELLER, "400")
    assert any("Недостатньо залишку" in t for t in s.texts())
    await send(dp, bot, SELLER, "◀️ Назад")
    assert "Оберіть товар" in s.texts()[-2]
    await send(dp, bot, SELLER, "❌ Скасувати")
    assert s.texts()[-1] == "Скасовано."
    assert db.one("SELECT COUNT(*) FROM sales")[0] == 0


@pytest.mark.asyncio
async def test_piece_product_sale(env):
    db, bot, dp, s = env
    pid = S.create_product(db, "Ravioli 300g", "pasta", "piece", Decimal("14"), piece_grams=300)
    S.create_purchase(db, ADMIN, "2026-09-01", "Pasta", [S.PurchaseLine(pid, 1500, Decimal("15.30"))])
    await send(dp, bot, SELLER, M_SALE)
    await click(dp, bot, SELLER, f"pp:sale:id:{pid}")
    await send(dp, bot, SELLER, "2")
    await click(dp, bot, SELLER, "sale:add")
    key = (await state_data(dp, bot, SELLER))["key"]
    await click(dp, bot, SELLER, f"sale:pay:cash:{key}")
    sale = db.one("SELECT * FROM sales")
    assert sale["total"] == "28.00"
    assert S.stock_of_product(db, pid) == 900


@pytest.mark.asyncio
async def test_seller_cannot_open_purchase_or_reports(env):
    db, bot, dp, s = env
    await send(dp, bot, SELLER, M_PURCHASE)
    assert "менеджеру" in s.texts()[-1]
    await send(dp, bot, SELLER, "📈 Звіти")
    assert "менеджеру" in s.texts()[-1]


@pytest.mark.asyncio
async def test_report_and_excel_export(env, tmp_path):
    db, bot, dp, s = env
    pid = S.create_product(db, "Speck", "meat", "weight", Decimal("26"))
    S.create_purchase(db, ADMIN, "2026-09-01", "A", [S.PurchaseLine(pid, 4336, Decimal("20.20"))])
    S.create_sale(db, SELLER, [S.SaleLine(pid, 3236, Decimal("26.10"))], "cash")
    await send(dp, bot, ADMIN, "📈 Звіти")
    await click(dp, bot, ADMIN, "rep:today")
    t = s.texts()[-2]
    assert "Виручка: <b>84,46 €</b>" in t and "Валовий прибуток" in t
    await click(dp, bot, ADMIN, "rep:xlsx:2026-01-01:2030-12-31")
    doc = [c for c in s.calls if type(c).__name__ == "SendDocument"][-1]
    from openpyxl import load_workbook
    wb = load_workbook(doc.document.path)
    assert wb.sheetnames == ["Товар", "Зведення", "Оплати", "Продажі за товарами", "Залишки", "Операції"]
    ws = wb["Товар"]
    row = [c.value for c in ws[3]]
    assert row[1] == "Speck" and row[8] == 4.336 and row[12] == 3.236 and row[14] == 84.46
