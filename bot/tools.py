"""Службові команди: python -m bot.tools <команда>

  seed-products products.csv   — масово створити товари з CSV (name;category;sale_mode;piece_grams;retail_price;sku)
  backup                       — зробити резервну копію у BACKUP_DIR
  restore file.db              — відновити базу з файлу
  add-admin <telegram_id>      — додати адміністратора
"""
from __future__ import annotations

import csv
import sys
from decimal import Decimal

from . import services as S
from .db import get_db


CAT_ALIASES = {"cheese": "cheese", "сир": "cheese", "сири": "cheese", "meat": "meat", "м'ясо": "meat",
               "м'ясні вироби": "meat", "pasta": "pasta", "паста": "pasta", "напівфабрикати": "pasta"}
MODE_ALIASES = {"weight": "weight", "вага": "weight", "на вагу": "weight", "кг": "weight",
                "piece": "piece", "шт": "piece", "поштучно": "piece", "штука": "piece"}


def _rows(name: str, data: bytes, sheet: str | None = None) -> list[dict]:
    """Читає CSV (роздільник ; або ,) чи XLSX (аркуш sheet, якщо є, інакше перший) у список словників за заголовками."""
    import io
    if name.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        text = data.decode("utf-8-sig", errors="replace")
        delim = ";" if text.splitlines()[0].count(";") >= text.splitlines()[0].count(",") else ","
        rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    if not rows:
        return []
    hdr = [str(h or "").strip().lower() for h in rows[0]]
    out = []
    for r in rows[1:]:
        if not r or all(c in (None, "") for c in r):
            continue
        d = {hdr[i]: (str(r[i]).strip() if i < len(r) and r[i] is not None else "") for i in range(len(hdr))}
        out.append(d)
    return out


def import_products(name: str, data: bytes) -> tuple[int, int, list[str]]:
    """-> (створено, пропущено як дублікати, помилки)"""
    db = get_db()
    created = skipped = 0
    errors: list[str] = []
    for i, row in enumerate(_rows(name, data), 2):
        nm = row.get("name", "")
        if not nm:
            continue
        try:
            cat = CAT_ALIASES[row.get("category", "").lower()]
            mode = MODE_ALIASES[row.get("sale_mode", "").lower() or ("piece" if row.get("piece_grams") else "weight")]
            pg = int(float(row["piece_grams"].replace(",", "."))) if row.get("piece_grams") else None
            price = Decimal((row.get("retail_price") or "0").replace(",", ".").replace(" ", ""))
        except (KeyError, ValueError) as e:
            errors.append(f"рядок {i} ({nm}): {e}")
            continue
        if db.one("SELECT 1 FROM products WHERE lower(name)=lower(?)", (nm,)):
            skipped += 1
            continue
        try:
            S.create_product(db, nm, cat, mode, price, pg, row.get("sku") or None)
            created += 1
        except Exception as e:
            errors.append(f"рядок {i} ({nm}): {e}")
    return created, skipped, errors


def import_opening_stock(name: str, data: bytes, user_id: int) -> tuple[int, list[str]]:
    """Файл з колонками name; kg; price_per_kg; expiry (дата, необов'язково). -> (створено партій, помилки)"""
    import datetime as dt
    db = get_db()
    created, errors = 0, []
    for i, row in enumerate(_rows(name, data), 2):
        nm = row.get("name", "")
        if not nm:
            continue
        prod = db.one("SELECT id FROM products WHERE lower(name)=lower(?)", (nm,))
        if not prod:
            errors.append(f"рядок {i}: товар «{nm}» не знайдено — спочатку імпортуйте товари")
            continue
        try:
            kg = Decimal(row["kg"].replace(",", ".").replace(" ", ""))
            price = Decimal(row["price_per_kg"].replace(",", ".").replace(" ", ""))
        except (KeyError, ValueError, ArithmeticError):
            errors.append(f"рядок {i} ({nm}): некоректні kg / price_per_kg")
            continue
        grams = int((kg * 1000).quantize(Decimal("1")))
        if grams <= 0:
            continue
        exp = None
        raw = row.get("expiry") or ""
        if raw:
            for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y"):
                try:
                    exp = dt.datetime.strptime(raw[:19], fmt).date().isoformat()
                    break
                except ValueError:
                    pass
        S.add_opening_stock(db, user_id, prod["id"], grams, price, exp, comment="імпорт початкових залишків")
        created += 1
    return created, errors


def seed_products(path: str) -> None:
    c, s, errs = import_products(path, open(path, "rb").read())
    print(f"створено: {c}, пропущено (вже є): {s}")
    for e in errs:
        print("помилка:", e)


def main(argv: list[str]) -> None:
    if not argv:
        print(__doc__)
        return
    cmd, *args = argv
    if cmd == "seed-products":
        seed_products(args[0])
    elif cmd == "backup":
        print("копія:", get_db().make_backup())
    elif cmd == "restore":
        get_db().restore_from(args[0])
        print("відновлено з", args[0])
    elif cmd == "add-admin":
        S.upsert_user(get_db(), int(args[0]), "admin")
        print("адміністратора додано")
    else:
        print(__doc__)


if __name__ == "__main__":
    main(sys.argv[1:])


# ======================= імпорт історії руху товарів =======================

HISTORY_COLS = ("period_start", "period_end", "name", "category", "sale_mode", "piece_grams", "retail_price",
                "buy_kg", "buy_eur", "sold_kg", "sold_eur", "writeoff_kg", "writeoff_eur", "closing_kg")


def _dec(v: str) -> Decimal:
    v = (v or "").replace(",", ".").replace(" ", "")
    return Decimal(v) if v not in ("", "-", "None") else Decimal(0)


def _date(v: str) -> str:
    import datetime as dt
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y"):
        try:
            return dt.datetime.strptime(v[:19], fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError(f"дата «{v}»")


def import_history(name: str, data: bytes, user_id: int) -> dict:
    """Імпорт руху товарів по періодах (аркуш «Рух товарів»).

    Для кожного рядка (період × товар):
      1) створює товар, якщо його ще немає (категорія, спосіб продажу, роздрібна ціна);
      2) закупівля buy_kg за buy_eur на дату period_start (від'ємна = повернення постачальнику → списання);
      3) продаж sold_kg за sold_eur на дату period_end (спосіб оплати «Інше», бо у звіті чеків немає);
      4) списання writeoff_kg на period_end (від'ємне = дооприбуткування);
      5) якщо після цього залишок ≠ closing_kg — інвентаризаційне коригування до closing_kg,
         щоб залишки в боті збігалися зі звітом рядок у рядок.
    Повторний імпорт того ж періоду захищений: якщо для товару вже є продаж з client_key цього періоду — рядок пропускається.
    """
    from .money import round_cents
    db = get_db()
    st = {"products_created": 0, "rows": 0, "skipped": 0, "aligned": 0, "errors": []}
    rows = _rows(name, data, "Рух товарів")
    for i, row in enumerate(rows, 2):
        nm = row.get("name", "")
        if not nm:
            continue
        try:
            p_start, p_end = _date(row["period_start"]), _date(row["period_end"])
            cat = CAT_ALIASES[row.get("category", "").lower()]
            mode = MODE_ALIASES.get((row.get("sale_mode") or "").lower(), "weight")
            pg = int(_dec(row.get("piece_grams", ""))) or None
            retail = _dec(row.get("retail_price", ""))
            buy_kg, buy_eur = _dec(row.get("buy_kg")), _dec(row.get("buy_eur"))
            sold_kg, sold_eur = _dec(row.get("sold_kg")), _dec(row.get("sold_eur"))
            wo_kg, wo_eur = _dec(row.get("writeoff_kg")), _dec(row.get("writeoff_eur"))
            closing_kg = _dec(row.get("closing_kg"))
        except (KeyError, ValueError, ArithmeticError) as e:
            st["errors"].append(f"рядок {i} ({nm}): {e}")
            continue
        prod = db.one("SELECT * FROM products WHERE lower(name)=lower(?)", (nm,))
        if not prod:
            if mode == "piece" and not pg:
                mode = "weight"
            pid = S.create_product(db, nm, cat, mode, retail if retail > 0 else Decimal(0), pg if mode == "piece" else None)
            prod = S.get_product(db, pid)
            st["products_created"] += 1
        elif retail > 0 and Decimal(prod["retail_price"]) == 0:
            S.update_product(db, prod["id"], retail_price=retail)
        pid = prod["id"]
        key = f"hist:{p_start}:{p_end}:{pid}"
        if db.one("SELECT 1 FROM sales WHERE client_key=?", (key,)):
            st["skipped"] += 1
            continue
        tag = f"імпорт історії {p_start}—{p_end}"
        g = lambda kg: int((kg * 1000).quantize(Decimal("1")))
        # 2) закупівля
        if buy_kg > 0:
            price = (buy_eur / buy_kg).quantize(Decimal("0.0001")) if buy_kg else Decimal(0)
            S.create_purchase(db, user_id, p_start, "імпорт історії",
                              [S.PurchaseLine(pid, g(buy_kg), price, None, p_start[:7], tag)], comment=tag)
        elif buy_kg < 0:
            _ensure_stock(db, user_id, pid, g(-buy_kg), (buy_eur / buy_kg) if buy_kg else Decimal(0), p_start, tag)
            S.write_off(db, user_id, pid, g(-buy_kg), f"повернення постачальнику ({tag})", op_date=p_start)
        # 3) продаж
        if sold_kg > 0:
            unit_cost = _last_price(db, pid)
            _ensure_stock(db, user_id, pid, g(sold_kg), unit_cost, p_start, tag)
            price_kg = (sold_eur / sold_kg).quantize(Decimal("0.0001"))
            S.create_sale(db, user_id, [S.SaleLine(pid, g(sold_kg), price_kg)], "other",
                          client_key=key, comment=tag, sold_at=p_end)
        else:
            # маркер періоду без продажу — щоб повторний імпорт пропускав рядок
            with db.tx() as c:
                c.execute("INSERT INTO sales(sold_at, sale_date, payment_method, status, total, cost_total, comment, "
                          "created_by, created_at, client_key) VALUES (?,?,?,?,?,?,?,?,?,?)",
                          (f"{p_end}T12:00:00+00:00", p_end, "other", "cancelled", "0", "0", tag + " (без продажу)",
                           user_id, p_end + "T12:00:00+00:00", key))
        # 4) списання / дооприбуткування
        if wo_kg > 0:
            _ensure_stock(db, user_id, pid, g(wo_kg), _last_price(db, pid), p_start, tag)
            S.write_off(db, user_id, pid, g(wo_kg), f"списання за звітом ({tag})", op_date=p_end)
        elif wo_kg < 0:
            _add_batch(db, user_id, pid, g(-wo_kg), (wo_eur / wo_kg) if wo_kg else _last_price(db, pid), p_end,
                       f"дооприбуткування за звітом ({tag})")
        # 5) вирівнювання до залишку на кінець періоду
        have = S.stock_of_product(db, pid)
        target = max(g(closing_kg), 0)
        if have != target:
            if target > have:
                _add_batch(db, user_id, pid, target - have, _last_price(db, pid), p_end, f"вирівнювання до звіту ({tag})")
            else:
                S.write_off(db, user_id, pid, have - target, f"вирівнювання до звіту ({tag})", op_date=p_end)
            st["aligned"] += 1
        st["rows"] += 1
    return st


def _last_price(db, pid: int) -> Decimal:
    b = db.one("SELECT landed_price_per_kg FROM batches WHERE product_id=? ORDER BY received_at DESC, id DESC LIMIT 1", (pid,))
    return Decimal(b["landed_price_per_kg"]) if b else Decimal(0)


def _add_batch(db, user_id, pid, grams, price_per_kg, date, comment):
    S.add_opening_stock(db, user_id, pid, grams, Decimal(price_per_kg).quantize(Decimal("0.0001")), None, date, comment)


def _ensure_stock(db, user_id, pid, need_grams, price_per_kg, date, tag):
    """Якщо у звіті продано/списано більше, ніж є в боті (розбіжності первинного файлу) — дооприбутковуємо різницю."""
    have = S.stock_of_product(db, pid)
    if have < need_grams:
        _add_batch(db, user_id, pid, need_grams - have, price_per_kg or Decimal(0), date, f"розбіжність звіту ({tag})")


# ======================= імпорт витрат =======================

EXP_TYPE_ALIASES = {"operating": "operating", "операційні": "operating", "goods": "goods", "оплата товару": "goods",
                    "tax": "tax", "податки": "tax", "investment": "investment", "інвестиції": "investment"}


def import_expenses(name: str, data: bytes, user_id: int) -> tuple[int, int, list[str]]:
    """Аркуш «Витрати»: date; type; category; amount; comment. -> (створено, пропущено-дублікати, помилки)"""
    db = get_db()
    created = skipped = 0
    errors = []
    for i, row in enumerate(_rows(name, data, "Витрати"), 2):
        if not row.get("category"):
            continue
        try:
            d_ = _date(row["date"])
            t = EXP_TYPE_ALIASES[row["type"].strip().lower()]
            amt = _dec(row["amount"])
        except (KeyError, ValueError, ArithmeticError) as e:
            errors.append(f"рядок {i}: {e}")
            continue
        if amt == 0:
            continue
        if db.one("SELECT 1 FROM expenses WHERE op_date=? AND exp_type=? AND category=? AND amount=? AND status='done'",
                  (d_, t, row["category"].strip(), str(amt.quantize(Decimal("0.01"))))):
            skipped += 1
            continue
        S.add_expense(db, user_id, d_, t, row["category"], amt, row.get("comment") or None)
        created += 1
    return created, skipped, errors
