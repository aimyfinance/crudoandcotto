"""Імпорт файлу інвентаризації (Excel): довнесення закупівель, прив'язки каси, виставлення фактичних залишків."""
from __future__ import annotations

import datetime as dt
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from . import services as S
from .db import get_db, now_utc, today_local
from .money import round_cents

SHEET_INV, SHEET_PUR, SHEET_ALIAS = "Інвентаризація", "Закупівля", "Прив'язки каси"


def _num(v) -> Decimal | None:
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return Decimal(str(v))
    s = str(v).replace("€", "").replace("кг", "").replace(" ", "").replace(",", ".")
    try:
        return Decimal(s)
    except Exception:
        return None


def _date(v) -> str:
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    s = str(v or "").strip()
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return today_local()


def _product(db, name: str):
    """Точна назва → підрядок → fuzzy (match_product)."""
    n = (name or "").strip()
    if not n:
        return None
    row = db.one("SELECT * FROM products WHERE lower(trim(name))=lower(?)", (n,))
    if row:
        return row
    cands = S.find_products(db, n)
    if len(cands) == 1:
        return cands[0]
    pid, how = S.match_product(db, n)
    return S.get_product(db, pid) if pid and how in ("exact", "alias", "auto") else None


def _rows(ws):
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if any(c not in (None, "") for c in r):
            yield i, list(r) + [None] * 8


def set_stock_to(db, user_id: int, pid: int, actual_grams: int, price_per_kg: Decimal | None, reason: str) -> dict:
    """Виставити залишок товару рівно actual_grams. Менше — списуємо з найстаріших партій; більше — додаємо до останньої
    партії (або нову за вказаною/останньою ціною). Партії з нульовою ціною переоцінюються, якщо дано ціну."""
    out = {"before": S.stock_of_product(db, pid), "after": actual_grams, "repriced": 0}
    if price_per_kg and price_per_kg > 0:
        with db.tx() as c:
            cur = c.execute("UPDATE batches SET price_per_kg=?, landed_price_per_kg=? WHERE product_id=? AND CAST(landed_price_per_kg AS REAL)=0",
                            (str(price_per_kg), str(price_per_kg), pid))
            out["repriced"] = cur.rowcount
            # собівартість уже проведених продажів із цих партій
            for slb in c.execute("SELECT slb.* FROM sale_line_batches slb JOIN batches b ON b.id=slb.batch_id WHERE b.product_id=? AND CAST(slb.cost AS REAL)=0", (pid,)).fetchall():
                cost = round_cents(Decimal(slb["grams"]) * price_per_kg / 1000)
                c.execute("UPDATE sale_line_batches SET cost=? WHERE id=?", (str(cost), slb["id"]))
                c.execute("UPDATE sale_lines SET cost=CAST(CAST(cost AS REAL)+? AS TEXT) WHERE id=?", (float(cost), slb["sale_line_id"]))
                c.execute("UPDATE sales SET cost_total=CAST(CAST(cost_total AS REAL)+? AS TEXT) WHERE id=(SELECT sale_id FROM sale_lines WHERE id=?)", (float(cost), slb["sale_line_id"]))
    delta = actual_grams - out["before"]
    if delta == 0:
        return out
    batches = S.batches_of_product(db, pid, only_open=True)
    if delta < 0:
        need = -delta
        for b in sorted(batches, key=lambda b: (b["received_at"], b["id"])):
            if need <= 0:
                break
            take = min(need, b["grams_left"])
            S.inventory_adjust(db, user_id, b["id"], b["grams_left"] - take, reason)
            need -= take
    else:
        if batches:
            last = sorted(batches, key=lambda b: (b["received_at"], b["id"]))[-1]
            S.inventory_adjust(db, user_id, last["id"], last["grams_left"] + delta, reason)
        else:
            lastp = db.one("SELECT landed_price_per_kg FROM batches WHERE product_id=? ORDER BY received_at DESC, id DESC LIMIT 1", (pid,))
            price = price_per_kg or (Decimal(lastp["landed_price_per_kg"]) if lastp else Decimal(0))
            S.add_opening_stock(db, user_id, pid, delta, price, None, today_local(), reason)
    return out


def import_inventory_xlsx(user_id: int, path: Path) -> dict:
    db = get_db()
    wb = load_workbook(path, data_only=True)
    res = {"purchases": 0, "purchase_lines": 0, "aliases": 0, "adjusted": [], "unchanged": 0, "errors": []}

    # 1) закупівлі (щоб продане наперед зарахувалось до правильних партій)
    if SHEET_PUR in wb.sheetnames:
        groups: dict[tuple, list] = {}
        for i, r in _rows(wb[SHEET_PUR]):
            date, supplier, name, kg, price, lot, expiry = r[:7]
            if not name or _num(kg) is None or _num(price) is None:
                continue
            p = _product(db, str(name))
            if not p:
                res["errors"].append(f"Закупівля, рядок {i}: товар «{name}» не знайдено"); continue
            exp = _date(expiry) if expiry not in (None, "") else None
            groups.setdefault((_date(date), str(supplier or "").strip()), []).append(
                S.PurchaseLine(p["id"], int(round(float(_num(kg)) * 1000)), round_cents(_num(price)), exp, str(lot).strip() if lot not in (None, "") else None, None))
        for (date, supplier), lines in groups.items():
            try:
                S.create_purchase(db, user_id, date, supplier, lines, Decimal(0), "імпорт інвентаризації", receive=True)
                res["purchases"] += 1; res["purchase_lines"] += len(lines)
            except Exception as e:
                res["errors"].append(f"Закупівля {supplier} {date}: {e}")

    # 2) прив'язки каси
    if SHEET_ALIAS in wb.sheetnames:
        for i, r in _rows(wb[SHEET_ALIAS]):
            raw, target = r[0], r[1]
            if not raw or not target:
                continue
            p = _product(db, str(target))
            if not p:
                res["errors"].append(f"Прив'язки, рядок {i}: товар «{target}» не знайдено"); continue
            S.set_alias(db, str(raw), p["id"], user_id); res["aliases"] += 1

    # 3) фактичні залишки
    if SHEET_INV in wb.sheetnames:
        for i, r in _rows(wb[SHEET_INV]):
            name, in_bot, expected, actual, price, comment = r[:6]
            if not name or _num(actual) is None:
                continue
            p = _product(db, str(name))
            if not p:
                res["errors"].append(f"Інвентаризація, рядок {i}: товар «{name}» не знайдено"); continue
            grams = int(round(float(_num(actual)) * 1000))
            try:
                o = set_stock_to(db, user_id, p["id"], grams, _num(price), f"інвентаризація {today_local()}" + (f": {comment}" if comment else ""))
                if o["before"] != o["after"] or o["repriced"]:
                    res["adjusted"].append(f"{p['name']}: {o['before'] / 1000:.3f} → {o['after'] / 1000:.3f} кг" + (f", переоцінено партій: {o['repriced']}" if o["repriced"] else ""))
                else:
                    res["unchanged"] += 1
            except Exception as e:
                res["errors"].append(f"Інвентаризація, рядок {i} ({p['name']}): {e}")
    with db.tx() as c:
        S.audit(c, user_id, "inventory.import", {k: (v if not isinstance(v, list) else len(v)) for k, v in res.items()})
    return res
