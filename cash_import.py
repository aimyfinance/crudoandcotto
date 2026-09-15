"""Імпорт звіту касового апарата для звірки з продажами бота.

Формат звіту каси залежить від моделі (у файлі витрат є «касовий апарат», але
його вивантаження ще не бачили). Тому імпортер універсальний:
  * приймає CSV (роздільник ; , або таб) та XLSX;
  * знаходить колонки за назвами заголовків (укр/нім/англ);
  * очікує або денні підсумки (дата, готівка, картка), або окремі чеки
    (дата, сума, спосіб оплати) — у другому випадку сам групує по днях.

Після отримання реального файлу з каси достатньо додати синоніми колонок у
словники нижче або написати окрему функцію-парсер.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import re
from decimal import Decimal, InvalidOperation

DATE_COLS = {"дата", "date", "datum", "день", "day", "tag", "beleg_datum", "belegdatum", "zeit", "datetime", "time"}
CASH_COLS = {"готівка", "cash", "bar", "bargeld", "barzahlung", "готівкою"}
CARD_COLS = {"картка", "card", "karte", "kartenzahlung", "bankomat", "kreditkarte", "картою", "безготівка"}
TOTAL_COLS = {"сума", "разом", "total", "summe", "betrag", "gesamt", "brutto", "amount", "umsatz", "всього"}
PAY_COLS = {"оплата", "спосіб оплати", "payment", "zahlung", "zahlungsart", "zahlungsmittel", "paymentmethod", "payment_method"}

CARD_WORDS = ("card", "karte", "картк", "bankomat", "kredit", "debit", "безгот")


def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip().lower())


def _money(v) -> Decimal:
    if v is None or v == "":
        return Decimal(0)
    if isinstance(v, (int, float)):
        return Decimal(str(v))
    s = str(v).strip().replace("€", "").replace("EUR", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")  # 1.234,56
    else:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        return Decimal(0)


def _date(v) -> str | None:
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v).strip()
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S",
                "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%d/%m/%Y %H:%M"):
        try:
            return dt.datetime.strptime(s[:19], fmt).date().isoformat()
        except ValueError:
            continue
    m = re.match(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", s)
    if m:
        return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
    return None


def _rows_from_bytes(name: str, data: bytes) -> list[list]:
    if name.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        return [list(r) for r in ws.iter_rows(values_only=True)]
    text = data.decode("utf-8-sig", errors="replace")
    sample = text[:2000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";"
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def parse_cash_report(name: str, data: bytes) -> dict[str, dict]:
    """-> {day: {"cash": Decimal, "card": Decimal}}"""
    rows = _rows_from_bytes(name, data)
    # знаходимо рядок-заголовок
    hdr_idx, cols = None, {}
    for i, row in enumerate(rows[:30]):
        names = [_norm(c) for c in row]
        found = {}
        for j, n in enumerate(names):
            if n in DATE_COLS and "date" not in found:
                found["date"] = j
            elif n in CASH_COLS:
                found["cash"] = j
            elif n in CARD_COLS:
                found["card"] = j
            elif n in TOTAL_COLS and "total" not in found:
                found["total"] = j
            elif n in PAY_COLS:
                found["pay"] = j
        if "date" in found and ({"cash", "card"} <= found.keys() or {"total", "pay"} <= found.keys() or "total" in found):
            hdr_idx, cols = i, found
            break
    if hdr_idx is None:
        raise ValueError(
            "Не вдалося розпізнати колонки. Потрібні заголовки: «Дата» + («Готівка», «Картка») "
            "або «Дата» + «Сума» + «Оплата». Надішліть файл мені — додам формат вашої каси."
        )
    out: dict[str, dict] = {}
    for row in rows[hdr_idx + 1:]:
        if not row or all(c in (None, "") for c in row):
            continue
        day = _date(row[cols["date"]] if cols["date"] < len(row) else None)
        if not day:
            continue
        e = out.setdefault(day, {"cash": Decimal(0), "card": Decimal(0)})
        if "cash" in cols and "card" in cols:
            e["cash"] += _money(row[cols["cash"]] if cols["cash"] < len(row) else 0)
            e["card"] += _money(row[cols["card"]] if cols["card"] < len(row) else 0)
        else:
            amt = _money(row[cols["total"]] if cols["total"] < len(row) else 0)
            pay = _norm(row[cols["pay"]]) if "pay" in cols and cols["pay"] < len(row) else ""
            if any(w in pay for w in CARD_WORDS):
                e["card"] += amt
            else:
                e["cash"] += amt
    if not out:
        raise ValueError("У файлі не знайдено жодного рядка з датою і сумою")
    return out
