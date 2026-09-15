"""Службові команди: python -m bot.tools <команда>

  seed-products products.csv   — масово створити товари з CSV (name;category;sale_mode;piece_grams;retail_price;sku)
  backup                       — зробити резервну копію у BACKUP_DIR
  restore file.db              — відновити базу з файлу
  add-admin <telegram_id>      — додати адміністратора
"""
from __future__ import annotations

import csv
import sys
from decimal import Decimal

from . import services as S
from .db import get_db


CAT_ALIASES = {"cheese": "cheese", "сир": "cheese", "сири": "cheese", "meat": "meat", "м'ясо": "meat",
               "м'ясні вироби": "meat", "pasta": "pasta", "паста": "pasta", "напівфабрикати": "pasta"}
MODE_ALIASES = {"weight": "weight", "вага": "weight", "на вагу": "weight", "кг": "weight",
                "piece": "piece", "шт": "piece", "поштучно": "piece", "штука": "piece"}


def _rows(name: str, data: bytes) -> list[dict]:
    """Читає CSV (роздільник ; або ,) чи XLSX з заголовками name;category;sale_mode;piece_grams;retail_price;sku."""
    import io
    if name.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        text = data.decode("utf-8-sig", errors="replace")
        delim = ";" if text.splitlines()[0].count(";") >= text.splitlines()[0].count(",") else ","
        rows = list(csv.reader(io.StringIO(text), delimiter=delim))
    if not rows:
        return []
    hdr = [str(h or "").strip().lower() for h in rows[0]]
    out = []
    for r in rows[1:]:
        if not r or all(c in (None, "") for c in r):
            continue
        d = {hdr[i]: (str(r[i]).strip() if i < len(r) and r[i] is not None else "") for i in range(len(hdr))}
        out.append(d)
    return out


def import_products(name: str, data: bytes) -> tuple[int, int, list[str]]:
    """-> (створено, пропущено як дублікати, помилки)"""
    db = get_db()
    created = skipped = 0
    errors: list[str] = []
    for i, row in enumerate(_rows(name, data), 2):
        nm = row.get("name", "")
        if not nm:
            continue
        try:
            cat = CAT_ALIASES[row.get("category", "").lower()]
            mode = MODE_ALIASES[row.get("sale_mode", "").lower() or ("piece" if row.get("piece_grams") else "weight")]
            pg = int(float(row["piece_grams"].replace(",", "."))) if row.get("piece_grams") else None
            price = Decimal((row.get("retail_price") or "0").replace(",", ".").replace(" ", ""))
        except (KeyError, ValueError) as e:
            errors.append(f"рядок {i} ({nm}): {e}")
            continue
        if db.one("SELECT 1 FROM products WHERE lower(name)=lower(?)", (nm,)):
            skipped += 1
            continue
        try:
            S.create_product(db, nm, cat, mode, price, pg, row.get("sku") or None)
            created += 1
        except Exception as e:
            errors.append(f"рядок {i} ({nm}): {e}")
    return created, skipped, errors


def import_opening_stock(name: str, data: bytes, user_id: int) -> tuple[int, list[str]]:
    """Файл з колонками name; kg; price_per_kg; expiry (дата, необов'язково). -> (створено партій, помилки)"""
    import datetime as dt
    db = get_db()
    created, errors = 0, []
    for i, row in enumerate(_rows(name, data), 2):
        nm = row.get("name", "")
        if not nm:
            continue
        prod = db.one("SELECT id FROM products WHERE lower(name)=lower(?)", (nm,))
        if not prod:
            errors.append(f"рядок {i}: товар «{nm}» не знайдено — спочатку імпортуйте товари")
            continue
        try:
            kg = Decimal(row["kg"].replace(",", ".").replace(" ", ""))
            price = Decimal(row["price_per_kg"].replace(",", ".").replace(" ", ""))
        except (KeyError, ValueError, ArithmeticError):
            errors.append(f"рядок {i} ({nm}): некоректні kg / price_per_kg")
            continue
        grams = int((kg * 1000).quantize(Decimal("1")))
        if grams <= 0:
            continue
        exp = None
        raw = row.get("expiry") or ""
        if raw:
            for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y"):
                try:
                    exp = dt.datetime.strptime(raw[:19], fmt).date().isoformat()
                    break
                except ValueError:
                    pass
        S.add_opening_stock(db, user_id, prod["id"], grams, price, exp, comment="імпорт початкових залишків")
        created += 1
    return created, errors


def seed_products(path: str) -> None:
    c, s, errs = import_products(path, open(path, "rb").read())
    print(f"створено: {c}, пропущено (вже є): {s}")
    for e in errs:
        print("помилка:", e)


def main(argv: list[str]) -> None:
    if not argv:
        print(__doc__)
        return
    cmd, *args = argv
    if cmd == "seed-products":
        seed_products(args[0])
    elif cmd == "backup":
        print("копія:", get_db().make_backup())
    elif cmd == "restore":
        get_db().restore_from(args[0])
        print("відновлено з", args[0])
    elif cmd == "add-admin":
        S.upsert_user(get_db(), int(args[0]), "admin")
        print("адміністратора додано")
    else:
        print(__doc__)


if __name__ == "__main__":
    main(sys.argv[1:])
